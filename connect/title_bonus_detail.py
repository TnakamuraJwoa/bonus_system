import logging
import math
from collections import defaultdict
from datetime import datetime, timedelta
from urllib.parse import urlencode

import openpyxl
from django.db import connections
from django.http import HttpResponse
from django.views import generic

from connect.models import MonthlyPeriod
from connect.sql.title_bonus_detail_sql import (
    TITLE_BONUS_DETAIL_CTE_SQL,
    TITLE_BONUS_DETAIL_SELECT_SQL,
    TITLE_BONUS_PURCHASE_INTRODUCER_TEAM_CTE_SQL,
    TITLE_BONUS_PURCHASE_DETAIL_SELECT_SQL,
    TITLE_BONUS_PURCHASE_PLACEMENT_TEAM_CTE_SQL,
)
from connect.templatetags.custom_filters import rank_label


logger = logging.getLogger(__name__)


class TitleBonusDetailView(generic.TemplateView):
    template_name = "title_bonus_detail.html"
    DEFAULT_PER_PAGE = 100
    MAX_PER_PAGE = 500
    PURCHASE_DETAIL_PER_PAGE = 500
    TEAM_TYPE_PLACEMENT = "placement"
    TEAM_TYPE_INTRODUCER = "introducer"
    TREE_MAX_DEPTH = 15
    TREE_MAX_NODES = 500
    TREE_SEARCH_MAX_DEPTH = 100

    def _get_per_page(self):
        try:
            per_page = int(self.request.GET.get("per_page") or self.DEFAULT_PER_PAGE)
        except ValueError:
            per_page = self.DEFAULT_PER_PAGE
        return max(1, min(per_page, self.MAX_PER_PAGE))

    def _get_page(self, total_pages):
        try:
            page = int(self.request.GET.get("page") or "1")
        except ValueError:
            page = 1
        return max(1, min(page, total_pages))

    def _get_purchase_page(self, total_pages):
        try:
            page = int(self.request.GET.get("purchase_page") or "1")
        except ValueError:
            page = 1
        return max(1, min(page, total_pages))

    @staticmethod
    def _build_base_qs(params):
        return urlencode({
            key: value
            for key, value in params.items()
            if value not in ("", None)
        })

    @staticmethod
    def _pagination_pages(current_page, total_pages, adjacent=3):
        if total_pages <= 1:
            return []

        pages = []
        if current_page > adjacent + 1:
            pages.append(1)
            pages.append(None)

        start = max(1, current_page - adjacent)
        end = min(total_pages, current_page + adjacent)
        pages.extend(range(start, end + 1))

        if current_page < total_pages - adjacent:
            pages.append(None)
            pages.append(total_pages)

        return pages

    @staticmethod
    def _empty_tree_context(unavailable_reason):
        return {
            "tree_ancestors": [],
            "tree_focus": None,
            "tree_children": [],
            "tree_truncated": False,
            "tree_unavailable_reason": unavailable_reason,
            "tree_node_count": 0,
            "tree_search_path_rows": [],
            "tree_search_target": None,
            "tree_search_not_found": False,
        }

    def _normalize_team_type(self):
        team_type = (self.request.GET.get("team_type") or self.TEAM_TYPE_INTRODUCER).strip()
        if team_type not in (self.TEAM_TYPE_PLACEMENT, self.TEAM_TYPE_INTRODUCER):
            return self.TEAM_TYPE_INTRODUCER
        return team_type

    def _get_purchase_team_cte_sql(self, team_type):
        if team_type == self.TEAM_TYPE_PLACEMENT:
            return TITLE_BONUS_PURCHASE_PLACEMENT_TEAM_CTE_SQL
        return TITLE_BONUS_PURCHASE_INTRODUCER_TEAM_CTE_SQL

    @staticmethod
    def _tree_parent_column(team_type):
        if team_type == TitleBonusDetailView.TEAM_TYPE_PLACEMENT:
            return "placement_code"
        return "introducer_code"

    @staticmethod
    def _tree_node_from_row(row):
        return {
            "jwoa_code": row.get("jwoa_code") or "",
            "send_bv_name": row.get("send_bv_name") or "",
            "rank": row.get("rank"),
            "placement_code": row.get("placement_code") or "",
            "introducer_code": row.get("introducer_code") or "",
            "rel_level": row.get("rel_level", 0),
            "children": [],
        }

    def _build_tree_children(self, parent_code, children_by_parent):
        nodes = []
        for row in children_by_parent.get(parent_code, []):
            node = self._tree_node_from_row(row)
            node["children"] = self._build_tree_children(
                node["jwoa_code"],
                children_by_parent,
            )
            nodes.append(node)
        return nodes

    def _execute_tree_rows(self, sql, params, log_message):
        with connections["rds"].cursor() as cursor:
            logger.info(log_message)
            cursor.execute(sql, params)
            columns = [col[0] for col in cursor.description]
            return [dict(zip(columns, row)) for row in cursor.fetchall()]

    def _execute_tree_scalar(self, sql, params):
        with connections["rds"].cursor() as cursor:
            cursor.execute(sql, params)
            row = cursor.fetchone()
            return int(row[0]) if row else 0

    def _build_users_tree_view(self, jwoa_code, team_type):
        result = {
            "tree_ancestors": [],
            "tree_focus": None,
            "tree_children": [],
            "tree_truncated": False,
            "tree_unavailable_reason": None,
            "tree_node_count": 0,
        }
        member_code = (jwoa_code or "").strip()
        if not member_code:
            result["tree_unavailable_reason"] = (
                "起点会員IDを入力して検索すると、Treeを表示できます。"
            )
            return result

        member_count = self._execute_tree_scalar(
            "SELECT COUNT(*) FROM bonus_db.users WHERE jmoa_code = %s",
            [member_code],
        )
        if member_count == 0:
            result["tree_unavailable_reason"] = (
                f"会員コード {member_code} のデータが見つかりません。"
            )
            return result
        if member_count > 1:
            result["tree_unavailable_reason"] = (
                f"会員コード {member_code} が複数件ヒットしています。"
                "完全一致で1件に特定できるコードを入力してください。"
            )
            return result

        parent_column = self._tree_parent_column(team_type)
        focus_sql = """
            SELECT
                u.jmoa_code AS jwoa_code,
                u.send_bv_name,
                u.`rank`,
                u.placement_code,
                u.introducer_code,
                0 AS rel_level
            FROM bonus_db.users AS u
            WHERE u.jmoa_code = %s
            LIMIT 1
        """
        focus_rows = self._execute_tree_rows(
            focus_sql,
            [member_code],
            "タイトルボーナス詳細 users Tree 起点SQLを実行します。",
        )
        if not focus_rows:
            result["tree_unavailable_reason"] = (
                f"会員コード {member_code} のデータが見つかりません。"
            )
            return result

        upline_sql = f"""
            WITH RECURSIVE upline AS (
                SELECT
                    u.jmoa_code AS jwoa_code,
                    u.send_bv_name,
                    u.`rank`,
                    u.placement_code,
                    u.introducer_code,
                    0 AS rel_level
                FROM bonus_db.users AS u
                WHERE u.jmoa_code = %s

                UNION ALL

                SELECT
                    u.jmoa_code AS jwoa_code,
                    u.send_bv_name,
                    u.`rank`,
                    u.placement_code,
                    u.introducer_code,
                    up.rel_level - 1 AS rel_level
                FROM bonus_db.users AS u
                JOIN upline AS up
                  ON u.jmoa_code = up.{parent_column}
                WHERE up.rel_level > -%s
                  AND up.{parent_column} IS NOT NULL
                  AND up.{parent_column} <> ''
            )
            SELECT
                jwoa_code,
                send_bv_name,
                `rank`,
                placement_code,
                introducer_code,
                rel_level
            FROM upline
            WHERE rel_level < 0
            ORDER BY rel_level ASC
            LIMIT %s
        """
        downline_sql = f"""
            WITH RECURSIVE downline AS (
                SELECT
                    u.jmoa_code AS jwoa_code,
                    u.send_bv_name,
                    u.`rank`,
                    u.placement_code,
                    u.introducer_code,
                    1 AS rel_level
                FROM bonus_db.users AS u
                WHERE u.{parent_column} = %s

                UNION ALL

                SELECT
                    u.jmoa_code AS jwoa_code,
                    u.send_bv_name,
                    u.`rank`,
                    u.placement_code,
                    u.introducer_code,
                    d.rel_level + 1 AS rel_level
                FROM bonus_db.users AS u
                JOIN downline AS d
                  ON u.{parent_column} = d.jwoa_code
                WHERE d.rel_level < %s
            )
            SELECT
                jwoa_code,
                send_bv_name,
                `rank`,
                placement_code,
                introducer_code,
                rel_level
            FROM downline
            ORDER BY rel_level ASC, jwoa_code ASC
            LIMIT %s
        """
        upline_rows = self._execute_tree_rows(
            upline_sql,
            [member_code, self.TREE_MAX_DEPTH, self.TREE_MAX_NODES],
            "タイトルボーナス詳細 users Tree 上位SQLを実行します。",
        )
        downline_rows = self._execute_tree_rows(
            downline_sql,
            [member_code, self.TREE_MAX_DEPTH, self.TREE_MAX_NODES],
            "タイトルボーナス詳細 users Tree 配下SQLを実行します。",
        )

        children_by_parent = defaultdict(list)
        for row in downline_rows:
            children_by_parent[row[parent_column]].append(row)

        result["tree_ancestors"] = [
            self._tree_node_from_row(row) for row in upline_rows
        ]
        result["tree_focus"] = self._tree_node_from_row(focus_rows[0])
        result["tree_children"] = self._build_tree_children(
            member_code,
            children_by_parent,
        )
        result["tree_node_count"] = (
            len(result["tree_ancestors"]) + 1 + len(downline_rows)
        )
        result["tree_truncated"] = (
            len(upline_rows) >= self.TREE_MAX_NODES
            or len(downline_rows) >= self.TREE_MAX_NODES
        )
        return result

    def _fetch_users_tree_search_path(self, root_code, tree_search, team_type):
        root = (root_code or "").strip()
        keyword = (tree_search or "").strip()
        if not root or not keyword:
            return []

        parent_column = self._tree_parent_column(team_type)
        code_prefix = f"{keyword}%"
        name_like = f"%{keyword}%"
        sql = f"""
            WITH RECURSIVE scope AS (
                SELECT
                    u.jmoa_code AS jwoa_code,
                    u.send_bv_name,
                    u.`rank`,
                    u.placement_code,
                    u.introducer_code,
                    0 AS rel_level,
                    CAST(u.jmoa_code AS CHAR(20000)) AS path_codes
                FROM bonus_db.users AS u
                WHERE u.jmoa_code = %s

                UNION ALL

                SELECT
                    u.jmoa_code AS jwoa_code,
                    u.send_bv_name,
                    u.`rank`,
                    u.placement_code,
                    u.introducer_code,
                    scope.rel_level + 1 AS rel_level,
                    CONCAT(scope.path_codes, ',', u.jmoa_code) AS path_codes
                FROM bonus_db.users AS u
                JOIN scope
                  ON u.{parent_column} = scope.jwoa_code
                WHERE scope.rel_level < %s
                  AND FIND_IN_SET(u.jmoa_code, scope.path_codes) = 0
            ),
            target AS (
                SELECT *
                FROM scope
                WHERE rel_level > 0
                  AND (
                      jwoa_code LIKE %s
                      OR send_bv_name LIKE %s
                  )
                ORDER BY
                    CASE
                        WHEN jwoa_code = %s THEN 0
                        WHEN jwoa_code LIKE %s THEN 1
                        ELSE 2
                    END,
                    rel_level,
                    jwoa_code
                LIMIT 1
            )
            SELECT
                s.jwoa_code,
                s.send_bv_name,
                s.`rank`,
                s.placement_code,
                s.introducer_code,
                s.rel_level,
                FIND_IN_SET(s.jwoa_code, target.path_codes) - 1 AS path_index,
                CASE WHEN s.jwoa_code = target.jwoa_code THEN 1 ELSE 0 END AS is_target
            FROM scope AS s
            JOIN target
              ON FIND_IN_SET(s.jwoa_code, target.path_codes) > 0
            ORDER BY path_index
        """
        return self._execute_tree_rows(
            sql,
            [
                root,
                self.TREE_SEARCH_MAX_DEPTH,
                code_prefix,
                name_like,
                keyword,
                code_prefix,
            ],
            "タイトルボーナス詳細 users Tree 経路SQLを実行します。",
        )

    def _build_purchase_tree_context(self, line_jwoa_code, tree_search, team_type):
        tree_context = self._build_users_tree_view(line_jwoa_code, team_type)
        tree_context.setdefault("tree_search_path_rows", [])
        tree_context.setdefault("tree_search_target", None)
        tree_context.setdefault("tree_search_not_found", False)

        if tree_search and not tree_context.get("tree_unavailable_reason"):
            tree_search_path_rows = self._fetch_users_tree_search_path(
                line_jwoa_code,
                tree_search,
                team_type,
            )
            tree_context["tree_search_path_rows"] = tree_search_path_rows
            tree_context["tree_search_target"] = (
                next((row for row in tree_search_path_rows if row.get("is_target")), None)
                if tree_search_path_rows
                else None
            )
            tree_context["tree_search_not_found"] = not tree_search_path_rows

        return tree_context

    def _build_purchase_team_context(self, team_type):
        if team_type == self.TEAM_TYPE_PLACEMENT:
            return {
                "purchase_team_title": "上位者チーム業績",
                "purchase_team_description": "起点会員IDの上位者Tree配下にあるタイトルボーナス対象購入を表示します。",
                "purchase_team_tree_title": "上位者チーム業績 Tree",
                "purchase_team_tree_description": "配下を上位者Treeで表示します。down会員IDまたはdown_nameを入力すると経路表示できます。",
                "purchase_line_header": "line_jwoa_code",
                "purchase_tree_ancestor_title": "上位者",
                "purchase_tree_ancestor_badge_label": "上位",
                "purchase_tree_direct_badge_label": "紹介者",
                "purchase_empty_message": "条件に一致する上位者チーム業績はありません。",
                "purchase_pagination_label": "上位者チーム業績ページ",
            }
        return {
            "purchase_team_title": "紹介者チーム業績",
            "purchase_team_description": "起点会員IDの紹介者Tree配下にあるタイトルボーナス対象購入を表示します。",
            "purchase_team_tree_title": "紹介者チーム業績 Tree",
            "purchase_team_tree_description": "配下を紹介者Treeで表示します。down会員IDまたはdown_nameを入力すると経路表示できます。",
            "purchase_line_header": "direct_introducer_line_code",
            "purchase_tree_ancestor_title": "紹介上位者",
            "purchase_tree_ancestor_badge_label": "紹介上位",
            "purchase_tree_direct_badge_label": "直紹介",
            "purchase_empty_message": "条件に一致する紹介者チーム業績はありません。",
            "purchase_pagination_label": "紹介者チーム業績ページ",
        }

    def _collect_purchase_tree_codes(self, tree_context):
        codes = []

        def append_code(node):
            code = node.get("jwoa_code")
            if code:
                codes.append(code)

        def walk(nodes):
            for node in nodes or []:
                append_code(node)
                walk(node.get("children") or [])

        for node in tree_context.get("tree_ancestors") or []:
            append_code(node)
        if tree_context.get("tree_focus"):
            append_code(tree_context["tree_focus"])
        walk(tree_context.get("tree_children") or [])
        for node in tree_context.get("tree_search_path_rows") or []:
            append_code(node)

        return list(dict.fromkeys(codes))

    def _fetch_tree_purchase_badges(self, period, member_codes):
        if not member_codes:
            return {}

        placeholders = ", ".join(["%s"] * len(member_codes))
        sql = f"""
            SELECT
                p.jwoa_code,
                p.order_type,
                CASE p.order_type
                    WHEN 101 THEN '再購入品'
                    WHEN 102 THEN '初回購入品'
                    WHEN 103 THEN 'ランクアップ購入品'
                    WHEN 105 THEN '特別対応購入品'
                    ELSE '対象外'
                END AS order_type_name,
                CASE p.order_type
                    WHEN 101 THEN 'badge-info'
                    WHEN 102 THEN 'badge-success'
                    WHEN 103 THEN 'badge-warning'
                    WHEN 105 THEN 'badge-danger'
                    ELSE 'badge-secondary'
                END AS badge_class,
                SUM(
                    CASE
                        WHEN p.order_type IN (101, 105)
                        THEN LEAST(IFNULL(p.bv, 0), 50)
                        WHEN p.order_type IN (102, 103)
                        THEN IFNULL(p.bv, 0)
                        ELSE 0
                    END
                ) AS bv_max50
            FROM bonus_db.purchase_info_list AS p
            WHERE p.jwoa_code IN ({placeholders})
              AND p.order_type IN (101, 102, 103, 105)
              AND p.register_year = %s
              AND p.register_month = %s
            GROUP BY
                p.jwoa_code,
                p.order_type
            HAVING bv_max50 > 0
            ORDER BY
                p.jwoa_code,
                p.order_type
        """
        params = [*member_codes, period.year, period.month]
        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 Tree購入バッジSQLを実行します。")
            cursor.execute(sql, params)
            columns = [col[0] for col in cursor.description]
            rows = [dict(zip(columns, row)) for row in cursor.fetchall()]

        badges = {}
        for row in rows:
            badges.setdefault(row["jwoa_code"], []).append(row)
        return badges

    def _apply_tree_purchase_badges(self, tree_context, badge_map):
        def apply_node(node):
            node["purchase_badges"] = badge_map.get(node.get("jwoa_code"), [])

        def walk(nodes):
            for node in nodes or []:
                apply_node(node)
                walk(node.get("children") or [])

        for node in tree_context.get("tree_ancestors") or []:
            apply_node(node)
        if tree_context.get("tree_focus"):
            apply_node(tree_context["tree_focus"])
        walk(tree_context.get("tree_children") or [])
        for node in tree_context.get("tree_search_path_rows") or []:
            apply_node(node)

    def _get_registered_periods(self):
        return (
            MonthlyPeriod.objects.using("rds")
            .all()
            .order_by("-year", "-month", "-kibetu")
        )

    @staticmethod
    def _build_calc_params(selected_kibetu, period):
        current_month_first = datetime(period.year, period.month, 1)
        prev_month_last = current_month_first - timedelta(days=1)

        return [
            period.year,
            period.month,
            selected_kibetu,
            period.year,
            period.month,
            prev_month_last.year,
            prev_month_last.month,
            period.year,
            period.month,
            prev_month_last.year,
            prev_month_last.month,
            period.year,
            period.month,
        ]

    @staticmethod
    def _build_filters(filters, table_alias):
        where = []
        params = []

        root_column = "m.root_jmoa_code" if table_alias == "m" else "t.root_jmoa_code"

        if filters["root_jwoa_code"]:
            where.append(f"{root_column} LIKE %s")
            params.append(f"%{filters['root_jwoa_code']}%")

        if filters["line_jwoa_code"]:
            where.append(f"{table_alias}.line_jwoa_code LIKE %s")
            params.append(f"%{filters['line_jwoa_code']}%")

        if filters["down_jwoa_code"]:
            where.append(f"{table_alias}.down_jwoa_code LIKE %s")
            params.append(f"%{filters['down_jwoa_code']}%")

        if table_alias == "m" and filters["match_level"]:
            where.append("m.match_level = %s")
            params.append(filters["match_level"])

        return where, params

    def _build_detail_sql(self, filters):
        where, params = self._build_filters(filters, "m")
        sql = TITLE_BONUS_DETAIL_SELECT_SQL
        if where:
            sql += "\n  AND " + "\n  AND ".join(where)
        return sql, params

    def _count_detail_rows(self, calc_params, filters):
        detail_sql, filter_params = self._build_detail_sql(filters)
        sql = f"""
            {TITLE_BONUS_DETAIL_CTE_SQL}
            SELECT COUNT(*)
            FROM (
                {detail_sql}
            ) AS detail_rows
        """
        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 発生明細COUNT SQLを実行します。")
            cursor.execute(sql, [*calc_params, *filter_params])
            row = cursor.fetchone()
            return int(row[0]) if row else 0

    def _fetch_detail_rows(self, calc_params, filters, limit, offset):
        detail_sql, filter_params = self._build_detail_sql(filters)
        sql = f"""
            {TITLE_BONUS_DETAIL_CTE_SQL}
            {detail_sql}
            ORDER BY
                m.root_jmoa_code,
                m.line_jwoa_code,
                m.match_level,
                m.tree_level,
                m.down_jwoa_code
            LIMIT %s OFFSET %s
        """
        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 発生明細SQLを実行します。")
            cursor.execute(sql, [*calc_params, *filter_params, limit, offset])
            columns = [col[0] for col in cursor.description]
            return [dict(zip(columns, row)) for row in cursor.fetchall()]

    def _fetch_line_summary_rows(self, calc_params, filters):
        detail_sql, filter_params = self._build_detail_sql(filters)
        sql = f"""
            {TITLE_BONUS_DETAIL_CTE_SQL}
            SELECT
                line_jwoa_code,
                line_name,
                COUNT(*) AS matched_count,
                MIN(match_level) AS min_match_level,
                MAX(match_level) AS max_match_level,
                MAX(sum_bv) AS line_sum_bv,
                SUM(bonus_amount) AS bonus_amount
            FROM (
                {detail_sql}
            ) AS detail_rows
            GROUP BY
                line_jwoa_code,
                line_name
            ORDER BY
                line_jwoa_code
        """
        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 ライン別サマリSQLを実行します。")
            cursor.execute(sql, [*calc_params, *filter_params])
            columns = [col[0] for col in cursor.description]
            return [dict(zip(columns, row)) for row in cursor.fetchall()]

    def _fetch_direct_introducer_summary_rows(self, calc_params, period, filters):
        detail_sql, filter_params = self._build_detail_sql(filters)
        sql = f"""
            {TITLE_BONUS_DETAIL_CTE_SQL}
            SELECT
                roots.root_jwoa_code,
                roots.root_name,
                u.jmoa_code AS introducer_jwoa_code,
                u.send_bv_name AS introducer_name,
                u.`rank`,
                CASE u.status_code
                    WHEN 1 THEN 'アクティブ'
                    WHEN 2 THEN '凍結'
                    WHEN 3 THEN '退会'
                    WHEN 4 THEN '中途解約'
                    WHEN 5 THEN '非アクティブ'
                    ELSE '-'
                END AS status_name,
                CASE u.status_code
                    WHEN 1 THEN 'badge-success'
                    WHEN 2 THEN 'badge-warning'
                    WHEN 3 THEN 'badge-danger'
                    WHEN 4 THEN 'badge-secondary'
                    WHEN 5 THEN 'badge-light text-muted'
                    ELSE 'badge-light text-muted'
                END AS status_badge_class,
                mt.title_id,
                COALESCE(tm.title_name, 'タイトルなし') AS title_name
            FROM (
                SELECT DISTINCT
                    root_jwoa_code,
                    root_name
                FROM (
                    {detail_sql}
                ) AS detail_rows
            ) AS roots
            JOIN bonus_db.users AS u
              ON u.introducer_code = roots.root_jwoa_code
            LEFT JOIN bonus_db.month_title AS mt
              ON mt.kibetu = %s
             AND mt.jwoa_code = u.jmoa_code
            LEFT JOIN bonus_db.title_master AS tm
              ON tm.title_id = mt.title_id
            ORDER BY
                roots.root_jwoa_code,
                mt.title_id DESC,
                u.jmoa_code
        """
        params = [
            *calc_params,
            *filter_params,
            period.kibetu,
        ]
        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 直紹介者サマリSQLを実行します。")
            cursor.execute(sql, params)
            columns = [col[0] for col in cursor.description]
            return [dict(zip(columns, row)) for row in cursor.fetchall()]

    def _fetch_three_star_summary_rows(self, calc_params, period, filters):
        current_month_first = datetime(period.year, period.month, 1)
        prev_month_last = current_month_first - timedelta(days=1)
        detail_sql, filter_params = self._build_detail_sql(filters)
        sql = f"""
            {TITLE_BONUS_DETAIL_CTE_SQL}
            , root_rows AS (
                SELECT DISTINCT
                    root_jwoa_code,
                    root_name
                FROM (
                    {detail_sql}
                ) AS detail_roots
            )
            , introducer_scope AS (
                SELECT
                    roots.root_jwoa_code,
                    roots.root_name,
                    u.jmoa_code,
                    u.send_bv_name,
                    u.introducer_code,
                    u.jmoa_code AS direct_introducer_jwoa_code,
                    u.send_bv_name AS direct_introducer_name,
                    1 AS introducer_level,
                    CASE WHEN ats.jwoa_code IS NOT NULL THEN 1 ELSE 0 END AS match_level,
                    CAST(u.jmoa_code AS CHAR(20000)) AS path_codes
                FROM root_rows AS roots
                JOIN bonus_db.users AS u
                  ON u.introducer_code = roots.root_jwoa_code
                LEFT JOIN active_three_star_dia AS ats
                  ON ats.jwoa_code = u.jmoa_code

                UNION ALL

                SELECT
                    scope.root_jwoa_code,
                    scope.root_name,
                    u.jmoa_code,
                    u.send_bv_name,
                    u.introducer_code,
                    scope.direct_introducer_jwoa_code,
                    scope.direct_introducer_name,
                    scope.introducer_level + 1 AS introducer_level,
                    scope.match_level +
                    CASE WHEN ats.jwoa_code IS NOT NULL THEN 1 ELSE 0 END AS match_level,
                    CONCAT(scope.path_codes, ',', u.jmoa_code) AS path_codes
                FROM introducer_scope AS scope
                JOIN bonus_db.users AS u
                  ON u.introducer_code = scope.jmoa_code
                LEFT JOIN active_three_star_dia AS ats
                  ON ats.jwoa_code = u.jmoa_code
                WHERE scope.introducer_level < 10000
                  AND scope.match_level < 6
                  AND FIND_IN_SET(u.jmoa_code, scope.path_codes) = 0
            )
            SELECT DISTINCT
                intro.root_jwoa_code,
                intro.root_name,
                intro.direct_introducer_jwoa_code AS line_jwoa_code,
                intro.direct_introducer_name AS line_name,
                intro.direct_introducer_jwoa_code,
                intro.direct_introducer_name,
                intro.introducer_level,
                intro.jmoa_code AS achiever_jwoa_code,
                intro.send_bv_name AS achiever_name,
                u.`rank`,
                ats.title_id AS achiever_title_id,
                COALESCE(tm.title_name, 'タイトルなし') AS achiever_title_name,
                intro.introducer_level AS tree_level,
                intro.match_level,
                CASE WHEN current_active.jwoa_code IS NOT NULL THEN 1 ELSE 0 END AS current_active_flg,
                CASE WHEN prev_active.jwoa_code IS NOT NULL THEN 1 ELSE 0 END AS prev_active_flg
            FROM introducer_scope AS intro
            JOIN active_three_star_dia AS ats
              ON ats.jwoa_code = intro.jmoa_code
            LEFT JOIN bonus_db.users AS u
              ON u.jmoa_code = intro.jmoa_code
            LEFT JOIN bonus_db.title_master AS tm
              ON tm.title_id = ats.title_id
            LEFT JOIN bonus_db.active_users AS current_active
              ON current_active.jwoa_code = intro.jmoa_code
             AND current_active.year = %s
             AND current_active.month = %s
             AND current_active.active_status = 1
            LEFT JOIN bonus_db.active_users AS prev_active
              ON prev_active.jwoa_code = intro.jmoa_code
             AND prev_active.year = %s
             AND prev_active.month = %s
             AND prev_active.active_status = 1
            ORDER BY
                intro.root_jwoa_code,
                intro.direct_introducer_jwoa_code,
                intro.match_level,
                intro.introducer_level,
                intro.jmoa_code
        """
        params = [
            *calc_params,
            *filter_params,
            period.year,
            period.month,
            prev_month_last.year,
            prev_month_last.month,
        ]
        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 3スター達成者サマリSQLを実行します。")
            cursor.execute(sql, params)
            columns = [col[0] for col in cursor.description]
            return [dict(zip(columns, row)) for row in cursor.fetchall()]

    def _build_purchase_detail_sql(self, filters):
        where = []
        filter_params = []
        sql = f"""
            {TITLE_BONUS_PURCHASE_DETAIL_SELECT_SQL}
        """
        purchase_where, purchase_params = self._build_purchase_filters(filters)
        where.extend(purchase_where)
        filter_params.extend(purchase_params)
        if where:
            sql += "\n  AND " + "\n  AND ".join(where)
        return sql, filter_params

    @staticmethod
    def _build_purchase_filters(filters):
        where = []
        params = []

        if filters["purchase_down_jwoa_code"]:
            where.append("t.down_jwoa_code LIKE %s")
            params.append(f"%{filters['purchase_down_jwoa_code']}%")

        if filters["purchase_down_name"]:
            where.append("t.down_name LIKE %s")
            params.append(f"%{filters['purchase_down_name']}%")

        if filters["purchase_title_id"]:
            if filters["purchase_title_id"] == "none":
                where.append("mt.title_id IS NULL")
            else:
                where.append("mt.title_id = %s")
                params.append(filters["purchase_title_id"])

        return where, params

    def _get_purchase_title_options(self, calc_params, period, filters):
        if not filters["purchase_line_jwoa_code"]:
            return []

        option_filters = {
            **filters,
            "purchase_title_id": "",
        }
        purchase_sql, filter_params = self._build_purchase_detail_sql(option_filters)
        team_cte_sql = self._get_purchase_team_cte_sql(filters["team_type"])
        sql = f"""
            {team_cte_sql}
            SELECT DISTINCT
                title_options.title_id,
                title_options.title_name
            FROM (
                SELECT
                    CASE
                        WHEN purchase_rows.title_name = 'タイトルなし'
                        THEN NULL
                        ELSE purchase_rows.title_id
                    END AS title_id,
                    purchase_rows.title_name
                FROM (
                    {purchase_sql}
                ) AS purchase_rows
            ) AS title_options
            ORDER BY
                CASE WHEN title_options.title_id IS NULL THEN 1 ELSE 0 END,
                title_options.title_id,
                title_options.title_name
        """
        params = [
            filters["purchase_line_jwoa_code"],
            period.year,
            period.month,
            period.kibetu,
            *filter_params,
        ]
        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 対象購入タイトル選択肢SQLを実行します。")
            cursor.execute(sql, params)
            columns = [col[0] for col in cursor.description]
            return [dict(zip(columns, row)) for row in cursor.fetchall()]

    def _count_purchase_detail_rows(self, calc_params, period, filters):
        if not filters["purchase_line_jwoa_code"]:
            return 0

        purchase_sql, filter_params = self._build_purchase_detail_sql(filters)
        team_cte_sql = self._get_purchase_team_cte_sql(filters["team_type"])
        sql = f"""
            {team_cte_sql}
            SELECT COUNT(*)
            FROM (
                {purchase_sql}
            ) AS purchase_rows
        """
        params = [
            filters["purchase_line_jwoa_code"],
            period.year,
            period.month,
            period.kibetu,
            *filter_params,
        ]
        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 チーム業績COUNT SQLを実行します。")
            cursor.execute(sql, params)
            row = cursor.fetchone()
            return int(row[0]) if row else 0

    def _fetch_purchase_summary(self, calc_params, period, filters):
        if not filters["purchase_line_jwoa_code"]:
            return None

        purchase_sql, filter_params = self._build_purchase_detail_sql(filters)
        team_cte_sql = self._get_purchase_team_cte_sql(filters["team_type"])
        sql = f"""
            {team_cte_sql}
            SELECT
                COUNT(*) AS row_count,
                COUNT(DISTINCT purchase_rows.down_jwoa_code) AS buyer_count,
                COALESCE(SUM(purchase_rows.original_bv), 0) AS original_bv_total,
                COALESCE(SUM(purchase_rows.bv_max50), 0) AS bv_max50_total,
                COALESCE(
                    SUM(purchase_rows.original_bv - purchase_rows.bv_max50),
                    0
                ) AS bv_limit_diff
            FROM (
                {purchase_sql}
            ) AS purchase_rows
        """
        params = [
            filters["purchase_line_jwoa_code"],
            period.year,
            period.month,
            period.kibetu,
            *filter_params,
        ]
        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 チーム業績サマリSQLを実行します。")
            cursor.execute(sql, params)
            columns = [col[0] for col in cursor.description]
            row = cursor.fetchone()
            return dict(zip(columns, row)) if row else None

    def _fetch_purchase_detail_rows(self, calc_params, period, filters, limit, offset):
        if not filters["purchase_line_jwoa_code"]:
            return []

        purchase_sql, filter_params = self._build_purchase_detail_sql(filters)
        team_cte_sql = self._get_purchase_team_cte_sql(filters["team_type"])
        sql = f"""
            {team_cte_sql}
            {purchase_sql}
        """
        sql += """
            ORDER BY
                t.root_jwoa_code,
                t.line_jwoa_code,
                t.tree_level,
                t.down_jwoa_code,
                p.bonus_payment_date,
                p.id
            LIMIT %s OFFSET %s
        """
        params = [
            filters["purchase_line_jwoa_code"],
            period.year,
            period.month,
            period.kibetu,
            *filter_params,
            limit,
            offset,
        ]
        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 チーム業績SQLを実行します。")
            cursor.execute(sql, params)
            columns = [col[0] for col in cursor.description]
            return [dict(zip(columns, row)) for row in cursor.fetchall()]

    def _fetch_registered_summary(self, selected_kibetu, filters):
        where = ["tbr.kibetu = %s"]
        params = [selected_kibetu]

        if filters["root_jwoa_code"]:
            where.append("tbr.root_jwoa_code LIKE %s")
            params.append(f"%{filters['root_jwoa_code']}%")

        if filters["down_jwoa_code"]:
            where.append("tbr.down_jwoa_code LIKE %s")
            params.append(f"%{filters['down_jwoa_code']}%")

        if filters["match_level"]:
            where.append("tbr.match_level = %s")
            params.append(filters["match_level"])

        sql = f"""
            SELECT
                COUNT(*) AS registered_count,
                COUNT(DISTINCT tbr.root_jwoa_code) AS root_count,
                SUM(tbr.bonus_amount) AS bonus_amount,
                MAX(tbr.created_at) AS latest_created_at
            FROM bonus_db.B_title_bonus_result AS tbr
            WHERE {" AND ".join(where)}
        """
        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 登録結果サマリSQLを実行します。")
            cursor.execute(sql, params)
            columns = [col[0] for col in cursor.description]
            row = cursor.fetchone()
            return dict(zip(columns, row)) if row else None

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)

        selected_kibetu = (self.request.GET.get("kibetu") or "").strip()
        purchase_kibetu = (
            self.request.GET.get("purchase_kibetu")
            or selected_kibetu
            or ""
        ).strip()
        team_type = self._normalize_team_type()
        filters = {
            "team_type": team_type,
            "root_jwoa_code": (self.request.GET.get("root_jwoa_code") or "").strip(),
            "line_jwoa_code": (self.request.GET.get("line_jwoa_code") or "").strip(),
            "down_jwoa_code": (self.request.GET.get("down_jwoa_code") or "").strip(),
            "match_level": (self.request.GET.get("match_level") or "").strip(),
            "purchase_down_jwoa_code": (
                self.request.GET.get("purchase_down_jwoa_code") or ""
            ).strip(),
            "purchase_down_name": (
                self.request.GET.get("purchase_down_name") or ""
            ).strip(),
            "purchase_line_jwoa_code": (
                self.request.GET.get("purchase_line_jwoa_code") or ""
            ).strip(),
            "purchase_title_id": (
                self.request.GET.get("purchase_title_id") or ""
            ).strip(),
        }
        purchase_tree_search = (self.request.GET.get("tree_search") or "").strip()
        if not purchase_tree_search:
            purchase_tree_search = (
                filters["purchase_down_jwoa_code"]
                or filters["purchase_down_name"]
            )
        per_page = self._get_per_page()

        ctx.update({
            "period_options": self._get_registered_periods(),
            "title_options": [],
            "selected_kibetu": selected_kibetu,
            "selected_period": None,
            "purchase_kibetu": purchase_kibetu,
            "purchase_period": None,
            "team_type": team_type,
            "team_type_placement": self.TEAM_TYPE_PLACEMENT,
            "team_type_introducer": self.TEAM_TYPE_INTRODUCER,
            "root_jwoa_code": filters["root_jwoa_code"],
            "line_jwoa_code": filters["line_jwoa_code"],
            "down_jwoa_code": filters["down_jwoa_code"],
            "match_level": filters["match_level"],
            "purchase_down_jwoa_code": filters["purchase_down_jwoa_code"],
            "purchase_down_name": filters["purchase_down_name"],
            "purchase_line_jwoa_code": filters["purchase_line_jwoa_code"],
            "purchase_title_id": filters["purchase_title_id"],
            "purchase_tree_search": purchase_tree_search,
            "purchase_tree": self._empty_tree_context(
                "起点会員IDを入力して「チーム業績を絞込」を押してください。"
            ),
            "match_level_options": ("1", "2", "3", "4", "5"),
            "per_page": per_page,
            "direct_introducer_summary_rows": [],
            "three_star_summary_rows": [],
            "detail_rows": [],
            "purchase_detail_rows": [],
            "purchase_summary": None,
            "purchase_detail_per_page": self.PURCHASE_DETAIL_PER_PAGE,
            "is_searched": bool(selected_kibetu),
        })
        ctx.update(self._build_purchase_team_context(team_type))

        base_params = {
            "kibetu": selected_kibetu,
            "purchase_kibetu": purchase_kibetu,
            "team_type": team_type,
            "root_jwoa_code": filters["root_jwoa_code"],
            "line_jwoa_code": filters["line_jwoa_code"],
            "down_jwoa_code": filters["down_jwoa_code"],
            "match_level": filters["match_level"],
            "purchase_down_jwoa_code": filters["purchase_down_jwoa_code"],
            "purchase_down_name": filters["purchase_down_name"],
            "purchase_line_jwoa_code": filters["purchase_line_jwoa_code"],
            "purchase_title_id": filters["purchase_title_id"],
            "tree_search": purchase_tree_search,
            "per_page": per_page,
        }
        purchase_base_params = {**base_params}
        placement_params = {
            **base_params,
            "team_type": self.TEAM_TYPE_PLACEMENT,
        }
        introducer_params = {
            **base_params,
            "team_type": self.TEAM_TYPE_INTRODUCER,
        }
        ctx["purchase_placement_qs"] = self._build_base_qs(placement_params)
        ctx["purchase_introducer_qs"] = self._build_base_qs(introducer_params)
        purchase_reset_params = {
            key: value
            for key, value in base_params.items()
            if key not in (
                "purchase_down_jwoa_code",
                "purchase_down_name",
                "purchase_line_jwoa_code",
                "purchase_title_id",
                "tree_search",
            )
        }
        ctx["purchase_reset_qs"] = self._build_base_qs(purchase_reset_params)
        if filters["purchase_line_jwoa_code"]:
            ctx["purchase_tree"] = self._build_purchase_tree_context(
                filters["purchase_line_jwoa_code"],
                purchase_tree_search,
                team_type,
            )

        if purchase_kibetu:
            purchase_period = (
                MonthlyPeriod.objects.using("rds")
                .filter(kibetu=purchase_kibetu)
                .first()
            )
            if purchase_period:
                ctx["purchase_period"] = purchase_period
                if filters["purchase_line_jwoa_code"]:
                    tree_codes = self._collect_purchase_tree_codes(ctx["purchase_tree"])
                    purchase_badge_map = self._fetch_tree_purchase_badges(
                        purchase_period,
                        tree_codes,
                    )
                    self._apply_tree_purchase_badges(
                        ctx["purchase_tree"],
                        purchase_badge_map,
                    )

                purchase_calc_params = self._build_calc_params(
                    purchase_kibetu,
                    purchase_period,
                )
                ctx["title_options"] = self._get_purchase_title_options(
                    purchase_calc_params,
                    purchase_period,
                    filters,
                )
                purchase_total_count = self._count_purchase_detail_rows(
                    purchase_calc_params,
                    purchase_period,
                    filters,
                )
                ctx["purchase_summary"] = self._fetch_purchase_summary(
                    purchase_calc_params,
                    purchase_period,
                    filters,
                )
                purchase_total_pages = max(
                    1,
                    math.ceil(purchase_total_count / self.PURCHASE_DETAIL_PER_PAGE),
                )
                purchase_page = self._get_purchase_page(purchase_total_pages)
                purchase_offset = (purchase_page - 1) * self.PURCHASE_DETAIL_PER_PAGE
                purchase_rows = []
                if purchase_total_count:
                    purchase_rows = self._fetch_purchase_detail_rows(
                        purchase_calc_params,
                        purchase_period,
                        filters,
                        self.PURCHASE_DETAIL_PER_PAGE,
                        purchase_offset,
                    )
                self._set_purchase_page_context(
                    ctx,
                    purchase_rows,
                    purchase_total_count,
                    purchase_total_pages,
                    purchase_page,
                    purchase_base_params,
                )
            else:
                self._set_purchase_page_context(
                    ctx,
                    [],
                    0,
                    1,
                    1,
                    purchase_base_params,
                )
        else:
            self._set_purchase_page_context(ctx, [], 0, 1, 1, purchase_base_params)

        if not selected_kibetu:
            return self._set_page_context(ctx, [], per_page, 0, 1, 1, base_params)

        period = MonthlyPeriod.objects.using("rds").filter(kibetu=selected_kibetu).first()
        if not period:
            return self._set_page_context(ctx, [], per_page, 0, 1, 1, base_params)

        ctx["selected_period"] = period
        calc_params = self._build_calc_params(selected_kibetu, period)
        ctx["direct_introducer_summary_rows"] = (
            self._fetch_direct_introducer_summary_rows(calc_params, period, filters)
        )
        ctx["three_star_summary_rows"] = self._fetch_three_star_summary_rows(
            calc_params,
            period,
            filters,
        )

        total_count = self._count_detail_rows(calc_params, filters)
        total_pages = max(1, math.ceil(total_count / per_page))
        page = self._get_page(total_pages)
        offset = (page - 1) * per_page
        rows = []
        if total_count:
            rows = self._fetch_detail_rows(calc_params, filters, per_page, offset)

        return self._set_page_context(
            ctx,
            rows,
            per_page,
            total_count,
            total_pages,
            page,
            base_params,
        )

    def _set_page_context(
        self,
        ctx,
        rows,
        per_page,
        total_count,
        total_pages,
        page,
        base_params,
    ):
        ctx["detail_rows"] = rows
        ctx["total_count"] = total_count
        ctx["per_page"] = per_page
        ctx["page"] = page
        ctx["total_pages"] = total_pages
        if total_count > 0:
            ctx["display_from"] = (page - 1) * per_page + 1
            ctx["display_to"] = min(ctx["display_from"] + len(rows) - 1, total_count)
        else:
            ctx["display_from"] = 0
            ctx["display_to"] = 0
        ctx["base_qs"] = self._build_base_qs(base_params)
        ctx["has_prev"] = page > 1
        ctx["has_next"] = page < total_pages
        ctx["prev_page"] = page - 1
        ctx["next_page"] = page + 1
        ctx["pagination_pages"] = self._pagination_pages(page, total_pages)
        return ctx

    def _set_purchase_page_context(
        self,
        ctx,
        rows,
        total_count,
        total_pages,
        page,
        base_params,
    ):
        ctx["purchase_detail_rows"] = rows
        ctx["purchase_total_count"] = total_count
        ctx["purchase_total_pages"] = total_pages
        ctx["purchase_page"] = page
        if total_count > 0:
            ctx["purchase_display_from"] = (
                (page - 1) * self.PURCHASE_DETAIL_PER_PAGE + 1
            )
            ctx["purchase_display_to"] = min(
                ctx["purchase_display_from"] + len(rows) - 1,
                total_count,
            )
        else:
            ctx["purchase_display_from"] = 0
            ctx["purchase_display_to"] = 0
        ctx["purchase_base_qs"] = self._build_base_qs(base_params)
        ctx["purchase_has_prev"] = page > 1
        ctx["purchase_has_next"] = page < total_pages
        ctx["purchase_prev_page"] = page - 1
        ctx["purchase_next_page"] = page + 1
        ctx["purchase_pagination_pages"] = self._pagination_pages(page, total_pages)
        return ctx


class TitleBonusDetailTreeExportView(TitleBonusDetailView):
    EXPORT_FETCH_SIZE = 2000

    def get(self, request, *args, **kwargs):
        root_code = (request.GET.get("purchase_line_jwoa_code") or "").strip()
        purchase_kibetu = (
            request.GET.get("purchase_kibetu")
            or request.GET.get("kibetu")
            or ""
        ).strip()

        wb = openpyxl.Workbook(write_only=True)
        ws = wb.create_sheet("紹介者チーム業績Tree")
        ws.append([
            "区分",
            "階層",
            "会員ID",
            "会員名",
            "rank",
            "紹介者ID",
            "紹介者名",
            "タイトルID",
            "タイトル",
            "order_type",
            "order_type名",
            "original_bv",
            "bv_max50",
        ])

        if not root_code or not purchase_kibetu:
            return self._excel_response(
                wb,
                "title_bonus_introducer_team_tree.xlsx",
            )

        period = MonthlyPeriod.objects.using("rds").filter(kibetu=purchase_kibetu).first()
        if not period:
            return self._excel_response(
                wb,
                "title_bonus_introducer_team_tree.xlsx",
            )

        sql = """
            WITH RECURSIVE ancestors AS (
                SELECT
                    u.jmoa_code AS jwoa_code,
                    u.send_bv_name,
                    u.`rank`,
                    u.introducer_code,
                    0 AS rel_level,
                    CAST(u.jmoa_code AS CHAR(20000)) AS path_codes
                FROM bonus_db.users AS u
                WHERE u.jmoa_code = %s

                UNION ALL

                SELECT
                    up.jmoa_code AS jwoa_code,
                    up.send_bv_name,
                    up.`rank`,
                    up.introducer_code,
                    a.rel_level - 1 AS rel_level,
                    CONCAT(up.jmoa_code, ',', a.path_codes) AS path_codes
                FROM ancestors AS a
                JOIN bonus_db.users AS up
                  ON up.jmoa_code = a.introducer_code
                WHERE a.rel_level > -%s
                  AND a.introducer_code IS NOT NULL
                  AND a.introducer_code <> ''
                  AND FIND_IN_SET(up.jmoa_code, a.path_codes) = 0
            ),
            descendants AS (
                SELECT
                    u.jmoa_code AS jwoa_code,
                    u.send_bv_name,
                    u.`rank`,
                    u.introducer_code,
                    0 AS rel_level,
                    CAST(u.jmoa_code AS CHAR(20000)) AS path_codes
                FROM bonus_db.users AS u
                WHERE u.jmoa_code = %s

                UNION ALL

                SELECT
                    child.jmoa_code AS jwoa_code,
                    child.send_bv_name,
                    child.`rank`,
                    child.introducer_code,
                    d.rel_level + 1 AS rel_level,
                    CONCAT(d.path_codes, ',', child.jmoa_code) AS path_codes
                FROM descendants AS d
                JOIN bonus_db.users AS child
                  ON child.introducer_code = d.jwoa_code
                WHERE d.rel_level < %s
                  AND FIND_IN_SET(child.jmoa_code, d.path_codes) = 0
            ),
            tree_nodes AS (
                SELECT
                    '紹介上位者' AS row_type,
                    jwoa_code,
                    send_bv_name,
                    `rank`,
                    introducer_code,
                    rel_level,
                    path_codes
                FROM ancestors
                WHERE rel_level < 0

                UNION ALL

                SELECT
                    CASE
                        WHEN rel_level = 0 THEN '起点'
                        ELSE '紹介配下'
                    END AS row_type,
                    jwoa_code,
                    send_bv_name,
                    `rank`,
                    introducer_code,
                    rel_level,
                    path_codes
                FROM descendants
            ),
            purchase_summary AS (
                SELECT
                    p.jwoa_code,
                    p.order_type,
                    CASE p.order_type
                        WHEN 101 THEN '再購入品'
                        WHEN 102 THEN '初回購入品'
                        WHEN 103 THEN 'ランクアップ購入品'
                        WHEN 105 THEN '特別対応購入品'
                        ELSE '対象外'
                    END AS order_type_name,
                    SUM(IFNULL(p.bv, 0)) AS original_bv,
                    CASE
                        WHEN p.order_type IN (101, 105)
                        THEN LEAST(SUM(IFNULL(p.bv, 0)), 50)
                        WHEN p.order_type IN (102, 103)
                        THEN SUM(IFNULL(p.bv, 0))
                        ELSE 0
                    END AS bv_max50
                FROM bonus_db.purchase_info_list AS p
                WHERE p.order_type IN (101, 102, 103, 105)
                  AND p.register_year = %s
                  AND p.register_month = %s
                GROUP BY
                    p.jwoa_code,
                    p.order_type
                HAVING bv_max50 > 0
            )
            SELECT
                n.row_type,
                n.rel_level,
                n.jwoa_code,
                n.send_bv_name,
                n.`rank`,
                n.introducer_code,
                parent.send_bv_name AS introducer_name,
                mt.title_id,
                COALESCE(tm.title_name, 'タイトルなし') AS title_name,
                ps.order_type,
                ps.order_type_name,
                ps.original_bv,
                ps.bv_max50,
                n.path_codes
            FROM tree_nodes AS n
            LEFT JOIN bonus_db.users AS parent
              ON parent.jmoa_code = n.introducer_code
            LEFT JOIN bonus_db.month_title AS mt
              ON mt.kibetu = %s
             AND mt.jwoa_code = n.jwoa_code
            LEFT JOIN bonus_db.title_master AS tm
              ON tm.title_id = mt.title_id
            LEFT JOIN purchase_summary AS ps
              ON ps.jwoa_code = n.jwoa_code
            ORDER BY
                CASE n.row_type
                    WHEN '紹介上位者' THEN 1
                    WHEN '起点' THEN 2
                    ELSE 3
                END,
                n.rel_level,
                n.path_codes,
                ps.order_type
        """
        params = [
            root_code,
            self.TREE_SEARCH_MAX_DEPTH,
            root_code,
            self.TREE_SEARCH_MAX_DEPTH,
            period.year,
            period.month,
            period.kibetu,
        ]

        with connections["rds"].cursor() as cursor:
            logger.info("タイトルボーナス詳細 紹介者チーム業績Tree Excel出力SQLを実行します。")
            cursor.execute(sql, params)
            while True:
                rows = cursor.fetchmany(self.EXPORT_FETCH_SIZE)
                if not rows:
                    break
                for row in rows:
                    ws.append([
                        row[0],
                        row[1],
                        row[2],
                        row[3],
                        rank_label(row[4]),
                        row[5],
                        row[6],
                        row[7],
                        row[8],
                        row[9],
                        row[10],
                        row[11] or 0,
                        row[12] or 0,
                    ])

        filename = f"title_bonus_introducer_team_tree_{purchase_kibetu}_{root_code}.xlsx"
        return self._excel_response(wb, filename)

    @staticmethod
    def _excel_response(workbook, filename):
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        workbook.save(response)
        return response
